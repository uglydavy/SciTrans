"""
Export translated documents to Microsoft Word format.
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available. Word export disabled.")


def export_to_word(
    translations: dict[str, str],
    source_pdf: str,
    output_path: str,
    source_lang: str = "en",
    target_lang: str = "fr",
) -> bool:
    """
    Export translations to Word document.

    Args:
        translations: Dict mapping block_id to translated text
        source_pdf: Path to source PDF
        output_path: Path for output Word document
        source_lang: Source language code
        target_lang: Target language code

    Returns:
        True if successful, False otherwise
    """
    if not DOCX_AVAILABLE:
        logger.error("python-docx not installed. Install with: pip install python-docx")
        return False

    try:
        doc = Document()

        # Title
        title = doc.add_heading(f"Translation: {source_lang} → {target_lang}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Add metadata
        doc.add_paragraph(f"Source: {Path(source_pdf).name}")
        doc.add_paragraph(f"Total blocks: {len(translations)}")
        doc.add_paragraph("")

        # Add translations
        for block_id, translated_text in translations.items():
            # Block header
            doc.add_heading(f"Block: {block_id}", level=2)

            # Translation text
            para = doc.add_paragraph(translated_text)
            para.style.font.size = Pt(11)

            doc.add_paragraph("")  # Spacing

        # Save
        doc.save(output_path)
        logger.info(f"Exported to Word: {output_path}")
        return True

    except Exception as e:
        logger.error(f"Failed to export to Word: {e}")
        return False
